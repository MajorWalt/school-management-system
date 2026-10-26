from django.contrib import messages
from django.db import transaction
from django.db.models import Max
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render

from accounts.utils import get_roles, is_admin
from core.activity import log_activity
from core.decorators import admin_or_teacher_required, admin_required, tenant_required
from scheduling.models import Form, Homeroom

from .forms import PromotionCriteriaForm, RequirementFormSet
from .models import AtRiskScore, PromotionCriteria
from .promotion import AREA_LABELS, evaluate_school
from .scoring import compute_school


def get_staff_profile(user):
    try:
        return user.staff_profile
    except Exception:
        return None


# ── At-risk score ─────────────────────────────────────────────────────────────


@admin_or_teacher_required
@tenant_required
def risk_list(request):
    school = request.school
    admin = is_admin(request.user, school)
    staff = get_staff_profile(request.user)
    roles = get_roles(request.user, school)

    if not admin and "teacher" not in roles:
        messages.error(request, "Access denied.")
        return redirect("portals:dashboard")

    # Teachers see only students in homerooms they are assigned to
    homerooms = None
    if not admin:
        homerooms = Homeroom.objects.filter(school=school, staff_members=staff) if staff else Homeroom.objects.none()
        if not homerooms.exists():
            messages.error(request, "Student Check-ins are only available to homeroom teachers.")
            return redirect("portals:dashboard")

    band = request.GET.get("band", "")
    form_pk = request.GET.get("form", "")

    latest = AtRiskScore.objects.filter(school=school).aggregate(m=Max("computed_on"))["m"]
    scores = AtRiskScore.objects.none()
    if latest:
        scores = AtRiskScore.objects.filter(school=school, computed_on=latest).select_related("student", "student__form", "student__homeroom")

    if homerooms is not None:
        scores = scores.filter(student__homeroom__in=homerooms)

    if band:
        scores = scores.filter(band=band)
    if form_pk:
        scores = scores.filter(student__form_id=form_pk)

    scores = scores.order_by("-score", "student__last_name")

    counts = {"high": 0, "watch": 0, "low": 0}
    for s in scores:
        counts[s.band] = counts.get(s.band, 0) + 1

    return render(
        request,
        "analytics/risk_list.html",
        {
            "scores": scores,
            "latest": latest,
            "counts": counts,
            "forms": Form.objects.filter(school=school),
            "band": band,
            "form_pk": form_pk,
            "is_admin": admin,
            "band_choices": AtRiskScore.BAND_CHOICES,
        },
    )


@admin_required
@tenant_required
def risk_recalculate(request):
    school = request.school
    if not is_admin(request.user, school):
        messages.error(request, "Access denied.")
        return redirect("analytics:risk_list")
    if request.method != "POST":
        return redirect("analytics:risk_list")

    r = compute_school(school)
    if r["error"]:
        messages.error(request, f"Scores not calculated: {r['error']}")
    else:
        log_activity(request, "risk_scores_recalculated", f"Recalculated at-risk scores: {r['scored']} students, {r['year']} Term {r['term']}.")
        msg = f"Scores recalculated for {r['scored']} students ({r['year']} Term {r['term']})."
        if r["skipped_no_grades"]:
            msg += f" {r['skipped_no_grades']} skipped with no grades this term."
        if r["skipped_small_form"]:
            msg += f" {r['skipped_small_form']} skipped in forms too small to compare."
        messages.success(request, msg)
    return redirect("analytics:risk_list")


# ── Promotion / graduation watch ──────────────────────────────────────────────


@admin_required
@tenant_required
def promotion_report(request):
    school = request.school
    if not is_admin(request.user, school):
        messages.error(request, "Access denied.")
        return redirect("portals:dashboard")

    form_pk = request.GET.get("form", "")
    outcome = request.GET.get("outcome", "")

    result = evaluate_school(school, form_id=form_pk or None, homerooms=None)
    rows = result["rows"]
    if outcome:
        rows = [r for r in rows if r["outcome"] == outcome]

    if request.GET.get("excel") == "1":
        return _promotion_excel(school, rows, result["year"])

    form_counts = {}
    for r in rows:
        form_counts[r["form"].name] = form_counts.get(r["form"].name, 0) + 1

    return render(
        request,
        "analytics/promotion_report.html",
        {
            "rows": rows,
            "result": result,
            "form_counts": form_counts,
            "forms": Form.objects.filter(school=school).order_by("order", "name"),
            "form_pk": form_pk,
            "outcome": outcome,
            "outcome_choices": PromotionCriteria.OUTCOME_CHOICES,
            "is_admin": True,
        },
    )


def _promotion_excel(school, rows, year):
    """Promotion watch table as an .xlsx download (same rows and filters as the page)."""
    import datetime
    import io

    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Promotion Watch"

    title_font = Font(name="Arial", bold=True, size=12)
    header_font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    header_fill = PatternFill("solid", start_color="1e40af")
    normal_font = Font(name="Arial", size=10)
    wrap = Alignment(wrap_text=True, vertical="top")

    year_name = year.name if year else ""
    ws.cell(row=1, column=1, value=f"Promotion Watch - {school.name} - {year_name}").font = title_font
    ws.cell(row=2, column=1, value=f"Generated {datetime.date.today():%B %d, %Y}. Based on year-to-date averages.").font = Font(name="Arial", size=9, italic=True)

    headers = [
        ("Student ID", 14),
        ("Student", 28),
        ("Form", 10),
        ("Homeroom", 14),
        ("Status", 20),
        ("Requirements not met", 60),
        ("Overall %", 11),
        ("Subjects below pass mark", 40),
    ]
    header_row = 4
    for col, (label, width) in enumerate(headers, start=1):
        c = ws.cell(row=header_row, column=col, value=label)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width = width

    for i, r in enumerate(rows, start=header_row + 1):
        unmet = "\n".join(f"{u['label']} ({u['need']}): {u['detail']}" for u in r["unmet"])
        failed = "\n".join(f"{name}: {avg:.1f}%" for name, avg in r["failed"])
        values = [
            getattr(r["student"], "student_id", "") or "",
            r["student"].get_full_name(),
            r["form"].name if r["form"] else "",
            r["homeroom"].name if r["homeroom"] else "",
            r["outcome_label"],
            unmet,
            r["overall"],
            failed,
        ]
        for col, v in enumerate(values, start=1):
            c = ws.cell(row=i, column=col, value=v)
            c.font = normal_font
            c.alignment = wrap
            if col == 7 and v is not None:
                c.number_format = "0.0"

    ws.freeze_panes = ws.cell(row=header_row + 1, column=3)
    if rows:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(headers))}{header_row + len(rows)}"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f"promotion_watch_{year_name.replace(' ', '_') or 'report'}.xlsx"
    response = HttpResponse(buf.read(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


@admin_required
@tenant_required
def criteria_list(request):
    school = request.school
    if not is_admin(request.user, school):
        messages.error(request, "Access denied.")
        return redirect("analytics:promotion_report")

    criteria = {c.form_id: c for c in PromotionCriteria.objects.filter(school=school).prefetch_related("requirements")}
    rows = []
    for f in Form.objects.filter(school=school).order_by("order", "name"):
        c = criteria.get(f.pk)
        reqs = []
        if c:
            for req in c.requirements.all():
                names = [AREA_LABELS.get(a, a) for a in (req.subject_areas or [])]
                reqs.append({"label": req.label, "subjects": ", ".join(names), "min_to_pass": req.min_to_pass, "count": len(names)})
        rows.append({"form": f, "criteria": c, "requirements": reqs})

    return render(request, "analytics/criteria_list.html", {"rows": rows, "is_admin": True})


@admin_required
@tenant_required
def criteria_edit(request, form_pk):
    school = request.school
    if not is_admin(request.user, school):
        messages.error(request, "Access denied.")
        return redirect("analytics:promotion_report")

    form_obj = get_object_or_404(Form, pk=form_pk, school=school)
    criteria = PromotionCriteria.objects.filter(school=school, form=form_obj).first()
    if criteria is None:
        last_form = Form.objects.filter(school=school).order_by("-order", "-name").first()
        criteria = PromotionCriteria(
            school=school,
            form=form_obj,
            outcome="graduation" if last_form and last_form.pk == form_obj.pk else "promotion",
        )

    cform = PromotionCriteriaForm(request.POST or None, instance=criteria)
    formset = RequirementFormSet(request.POST or None, instance=criteria, prefix="req")

    if request.method == "POST" and cform.is_valid() and formset.is_valid():
        with transaction.atomic():
            saved = cform.save(commit=False)
            saved.school = school
            saved.form = form_obj
            saved.updated_by = request.user
            saved.save()
            formset.instance = saved
            formset.save()
        log_activity(request, "promotion_criteria_saved", f"Saved promotion criteria for {form_obj.name}.")
        messages.success(request, f"Criteria for {form_obj.name} saved.")
        return redirect("analytics:criteria_list")

    return render(
        request,
        "analytics/criteria_form.html",
        {
            "form_obj": form_obj,
            "criteria": criteria,
            "cform": cform,
            "formset": formset,
            "is_admin": True,
        },
    )


@admin_required
@tenant_required
def criteria_delete(request, form_pk):
    school = request.school
    if not is_admin(request.user, school):
        messages.error(request, "Access denied.")
        return redirect("analytics:promotion_report")
    if request.method != "POST":
        return redirect("analytics:criteria_edit", form_pk=form_pk)

    criteria = get_object_or_404(PromotionCriteria, school=school, form_id=form_pk)
    name = criteria.form.name
    criteria.delete()
    log_activity(request, "promotion_criteria_deleted", f"Removed promotion criteria for {name}.")
    messages.warning(request, f"Criteria for {name} removed.")
    return redirect("analytics:criteria_list")