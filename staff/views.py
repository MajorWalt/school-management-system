import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.paginator import Paginator
from django.db import transaction, IntegrityError
from django.db.models import Q, Count
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.views.decorators.http import require_http_methods
from core.decorators import tenant_required, admin_required
from core.activity import log_activity
from scheduling.models import AcademicYear, TermConfig, SchoolEvent
from .forms import StaffForm, StaffClockInForm
from .models import Staff, StaffClockIn


@admin_required
def staff_list(request):
    staff = Staff.objects.filter(school=request.school)
    query = request.GET.get("q", "")
    if query:
        staff = staff.filter(first_name__icontains=query) | staff.filter(last_name__icontains=query) | staff.filter(employee_number__icontains=query)
    return render(request, "staff/staff_list.html", {"staff": staff, "query": query})
    pass


@admin_required
def staff_detail(request, pk):
    member = get_object_or_404(Staff, pk=pk, school=request.school)
    return render(request, "staff/staff_detail.html", {"member": member})
    pass


@admin_required
def staff_add(request):
    form = StaffForm(request.POST or None, school=request.school)
    if request.method == "POST" and form.is_valid():
        staff = form.save(commit=False)
        staff.school = request.school
        staff.save()
        log_activity(request, "staff_add", f"Added staff member {staff.get_full_name()} ({staff.employee_number}).")
        messages.success(request, f"{staff.get_full_name()} added successfully.")
        return redirect("staff:list")
    return render(request, "staff/staff_form.html", {"form": form, "title": "Add Staff Member"})
    pass


@admin_required
def staff_edit(request, pk):
    member = get_object_or_404(Staff, pk=pk, school=request.school)
    form = StaffForm(request.POST or None, instance=member, school=request.school)
    if request.method == "POST" and form.is_valid():
        form.save()
        log_activity(request, "staff_edit", f"Edited staff member {member.get_full_name()} ({member.employee_number}).")
        messages.success(request, f"{member.get_full_name()} updated successfully.")
        return redirect("staff:detail", pk=pk)
    return render(request, "staff/staff_form.html", {"form": form, "title": "Edit Staff Member"})
    pass


@admin_required
def staff_deactivate(request, pk):
    member = get_object_or_404(Staff, pk=pk, school=request.school)
    member.active = False
    member.save()
    log_activity(request, "staff_deactivate", f"Deactivated staff member {member.get_full_name()} ({member.employee_number}).")
    messages.warning(request, f"{member.get_full_name()} has been deactivated.")
    return redirect("staff:list")
    pass


@admin_required
def staff_reactivate(request, pk):
    member = get_object_or_404(Staff, pk=pk, school=request.school)
    member.active = True
    member.save()
    log_activity(request, "staff_reactivate", f"Reactivated staff member {member.get_full_name()} ({member.employee_number}).")
    messages.success(request, f"{member.get_full_name()} has been reactivated.")
    return redirect("staff:list")
    pass


# ============================================================================
# CLOCK-IN VIEWS
# ============================================================================


def _is_school_day(school, date):
    """Check if date is a school day (not Sat/Sun or marked as non-school day)."""
    # Check if it's a weekend
    if date.weekday() in (5, 6):  # Saturday=5, Sunday=6
        return False
        pass

    # Check if there's a SchoolEvent marking this as a non-school day
    non_school_event = SchoolEvent.objects.filter(school=school, date=date, is_school_day=False).exists()

    return not non_school_event
    pass


@login_required
@tenant_required
@require_http_methods(["POST"])
def clockin_self(request):
    """Staff member clocks in for today."""
    school = request.school

    # Check if user has a staff profile
    try:
        staff = Staff.objects.get(school=school, user=request.user)
    except Staff.DoesNotExist:
        return HttpResponse("Not authorized", status=403)
        pass

    today = timezone.localtime().date()

    # Check if it's a school day
    if not _is_school_day(school, today):
        return HttpResponse("Cannot clock in on non-school days", status=400)
        pass

    # Get location from POST data
    latitude = request.POST.get("latitude")
    longitude = request.POST.get("longitude")

    # Get or create clock-in entry (allow multiple clicks without error)
    with transaction.atomic():
        clock_in, created = StaffClockIn.objects.get_or_create(
            school=school,
            staff=staff,
            date=today,
            defaults={
                "clock_in_time": timezone.localtime().time(),
                "status": "present",
                "recorded_by": request.user,
                "latitude": latitude if latitude else None,
                "longitude": longitude if longitude else None,
            },
        )

        # Calculate distance if location provided
        if created and latitude and longitude:
            clock_in.calculate_distance_and_status()
            clock_in.save()
            pass
        pass

    if created:
        log_activity(request, "other", f"Clocked in at {clock_in.clock_in_time.strftime('%H:%M')}")
        pass

    # Return the updated card partial
    return render(
        request,
        "staff/partials/clockin_card.html",
        {
            "clock_in": clock_in,
            "staff": staff,
            "today": today,
            "is_school_day": True,
        },
    )
    pass

    today = timezone.localtime().date()

    # Check if it's a school day
    if not _is_school_day(school, today):
        return HttpResponse("Cannot clock in on non-school days", status=400)
        pass

    # Get or create clock-in entry (allow multiple clicks without error)
    with transaction.atomic():
        clock_in, created = StaffClockIn.objects.get_or_create(
            school=school,
            staff=staff,
            date=today,
            defaults={
                "clock_in_time": timezone.localtime().time(),
                "status": "present",
                "recorded_by": request.user,
            },
        )
        pass

    if created:
        log_activity(request, "other", f"Clocked in at {clock_in.clock_in_time.strftime('%H:%M')}")
        pass

    # Return the updated card partial
    return render(
        request,
        "staff/partials/clockin_card.html",
        {
            "clock_in": clock_in,
            "staff": staff,
            "today": today,
            "is_school_day": True,
        },
    )
    pass


@login_required
@tenant_required
def clockin_list(request):
    """List clock-in records. Admins see all, staff see only their own."""
    school = request.school

    # Check if user is admin
    from accounts.models import UserRole

    is_admin = UserRole.objects.filter(user=request.user, school=school, role="admin").exists() or request.user.is_superuser

    # Get staff profile for non-admin users
    staff_profile = None
    if not is_admin:
        try:
            staff_profile = Staff.objects.get(school=school, user=request.user)
        except Staff.DoesNotExist:
            return HttpResponse("Not authorized", status=403)
            pass
        pass

    # Base queryset
    if is_admin:
        clock_ins = StaffClockIn.objects.filter(school=school).select_related("staff", "recorded_by")
    else:
        clock_ins = StaffClockIn.objects.filter(school=school, staff=staff_profile).select_related("staff", "recorded_by")
        pass

    # Get current academic year and term for default filter
    current_year = AcademicYear.objects.filter(school=school, is_current=True).first()
    current_term = None
    term_start = None
    term_end = None

    if current_year:
        today = datetime.date.today()
        current_term = TermConfig.objects.filter(academic_year=current_year, start_date__lte=today, end_date__gte=today).first()

        if current_term:
            term_start = current_term.start_date
            term_end = current_term.end_date
        else:
            # Use academic year dates if no current term
            term_start = current_year.start_date
            term_end = current_year.end_date
            pass
        pass

    # Filters
    filter_year = request.GET.get("year", str(current_year.pk) if current_year else "")
    filter_term = request.GET.get("term", str(current_term.term_number) if current_term else "")
    filter_month = request.GET.get("month", "")
    filter_staff = request.GET.get("staff", "")
    filter_status = request.GET.get("status", "")
    filter_date_from = request.GET.get("date_from", "")
    filter_date_to = request.GET.get("date_to", "")

    # Apply filters
    if filter_year:
        try:
            year = AcademicYear.objects.get(pk=filter_year, school=school)
            if filter_term:
                term = TermConfig.objects.filter(academic_year=year, term_number=filter_term).first()
                if term:
                    clock_ins = clock_ins.filter(date__gte=term.start_date, date__lte=term.end_date)
                    term_start = term.start_date
                    term_end = term.end_date
                    pass
            else:
                clock_ins = clock_ins.filter(date__gte=year.start_date, date__lte=year.end_date)
                term_start = year.start_date
                term_end = year.end_date
                pass
        except (AcademicYear.DoesNotExist, ValueError):
            pass
        pass

    if is_admin and filter_staff:
        clock_ins = clock_ins.filter(staff__pk=filter_staff)
        pass

    if filter_status:
        clock_ins = clock_ins.filter(status=filter_status)
        pass

    if filter_month:
        try:
            # Format: YYYY-MM
            year_val, month_val = filter_month.split("-")
            clock_ins = clock_ins.filter(date__year=year_val, date__month=month_val)
        except (ValueError, IndexError):
            pass
        pass

    if filter_date_from:
        clock_ins = clock_ins.filter(date__gte=filter_date_from)
        pass

    if filter_date_to:
        clock_ins = clock_ins.filter(date__lte=filter_date_to)
        pass

    # Calculate summary stats for filtered set
    summary = clock_ins.aggregate(
        present=Count("id", filter=Q(status="present")),
        sick=Count("id", filter=Q(status="sick")),
        excused=Count("id", filter=Q(status="excused")),
    )

    # Pagination
    paginator = Paginator(clock_ins, 50)
    page_obj = paginator.get_page(request.GET.get("page", 1))

    # Get academic years and staff for filters (admins only)
    academic_years = AcademicYear.objects.filter(school=school).order_by("-start_date")
    staff_list = Staff.objects.filter(school=school, active=True).order_by("last_name", "first_name") if is_admin else None
    terms = TermConfig.objects.filter(academic_year=current_year).order_by("term_number") if current_year else None

    return render(
        request,
        "staff/clockin_list.html",
        {
            "page_obj": page_obj,
            "summary": summary,
            "is_admin": is_admin,
            "academic_years": academic_years,
            "staff_list": staff_list,
            "terms": terms,
            "current_year": current_year,
            "current_term": current_term,
            "filter_year": filter_year,
            "filter_term": filter_term,
            "filter_month": filter_month,
            "filter_staff": filter_staff,
            "filter_status": filter_status,
            "filter_date_from": filter_date_from,
            "filter_date_to": filter_date_to,
            "status_choices": StaffClockIn.STATUS_CHOICES,
        },
    )
    pass


@admin_required
def clockin_create(request):
    """Admin creates a clock-in record."""
    school = request.school
    form = StaffClockInForm(request.POST or None, school=school)

    if request.method == "POST" and form.is_valid():
        try:
            with transaction.atomic():
                clock_in = form.save(commit=False)
                clock_in.school = school
                clock_in.recorded_by = request.user
                clock_in.full_clean()
                clock_in.save()
                pass

            log_activity(request, "other", f"Created clock-in for {clock_in.staff.get_display_name()} on {clock_in.date} ({clock_in.get_status_display()})")
            messages.success(request, f"Clock-in record created for {clock_in.staff.get_display_name()}.")
            return redirect("staff:clockin_list")
        except IntegrityError:
            form.add_error(None, "A clock-in record already exists for this staff member on this date.")
            pass
        pass

    return render(
        request,
        "staff/clockin_form.html",
        {
            "form": form,
            "title": "Create Clock-In Record",
        },
    )
    pass


@admin_required
def clockin_edit(request, pk):
    """Admin edits a clock-in record."""
    school = request.school
    clock_in = get_object_or_404(StaffClockIn, pk=pk, school=school)
    form = StaffClockInForm(request.POST or None, instance=clock_in, school=school)

    if request.method == "POST" and form.is_valid():
        try:
            with transaction.atomic():
                updated = form.save(commit=False)
                updated.full_clean()
                updated.save()
                pass

            log_activity(request, "other", f"Edited clock-in for {clock_in.staff.get_display_name()} on {clock_in.date}")
            messages.success(request, "Clock-in record updated successfully.")
            return redirect("staff:clockin_list")
        except IntegrityError:
            form.add_error(None, "A clock-in record already exists for this staff member on this date.")
            pass
        pass

    return render(
        request,
        "staff/clockin_form.html",
        {
            "form": form,
            "title": "Edit Clock-In Record",
            "clock_in": clock_in,
        },
    )
    pass


@admin_required
@require_http_methods(["GET", "POST"])
def clockin_delete(request, pk):
    """Admin deletes a clock-in record."""
    school = request.school
    clock_in = get_object_or_404(StaffClockIn, pk=pk, school=school)

    if request.method == "POST":
        staff_name = clock_in.staff.get_display_name()
        date = clock_in.date
        clock_in.delete()

        log_activity(request, "other", f"Deleted clock-in for {staff_name} on {date}")
        messages.warning(request, f"Clock-in record for {staff_name} on {date} has been deleted.")
        return redirect("staff:clockin_list")
        pass

    return render(
        request,
        "staff/clockin_confirm_delete.html",
        {
            "clock_in": clock_in,
        },
    )
    pass


@admin_required
def clockin_export_excel(request):
    """Export clock-in records to Excel with filters."""
    school = request.school

    # Get filter parameters
    filter_year = request.GET.get("year", "")
    filter_term = request.GET.get("term", "")
    filter_month = request.GET.get("month", "")
    filter_staff = request.GET.get("staff", "")
    filter_status = request.GET.get("status", "")

    # Base queryset
    clock_ins = StaffClockIn.objects.filter(school=school).select_related("staff", "recorded_by").order_by("-date", "-clock_in_time")

    # Apply filters
    if filter_year:
        try:
            year = AcademicYear.objects.get(pk=filter_year, school=school)
            if filter_term:
                term = TermConfig.objects.filter(academic_year=year, term_number=filter_term).first()
                if term:
                    clock_ins = clock_ins.filter(date__gte=term.start_date, date__lte=term.end_date)
                    pass
            else:
                clock_ins = clock_ins.filter(date__gte=year.start_date, date__lte=year.end_date)
                pass
        except (AcademicYear.DoesNotExist, ValueError):
            pass
        pass

    if filter_month:
        try:
            # Format: YYYY-MM
            year_val, month_val = filter_month.split("-")
            clock_ins = clock_ins.filter(date__year=year_val, date__month=month_val)
        except (ValueError, IndexError):
            pass
        pass

    if filter_staff:
        clock_ins = clock_ins.filter(staff__pk=filter_staff)
        pass

    if filter_status:
        clock_ins = clock_ins.filter(status=filter_status)
        pass

    # Create workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Clock-Ins"

    # Define styles
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")

    thin_border = Border(left=Side(style="thin"), right=Side(style="thin"), top=Side(style="thin"), bottom=Side(style="thin"))

    # Headers
    headers = ["Date", "Staff Name", "Employee #", "Time", "Status", "Note", "Recorded By", "Latitude", "Longitude", "Distance (m)", "Within Radius"]

    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border
        pass

    # Data rows
    row_num = 2
    for clock_in in clock_ins:
        ws.cell(row=row_num, column=1, value=clock_in.date.strftime("%Y-%m-%d")).border = thin_border
        ws.cell(row=row_num, column=2, value=clock_in.staff.get_display_name()).border = thin_border
        ws.cell(row=row_num, column=3, value=clock_in.staff.employee_number).border = thin_border
        ws.cell(row=row_num, column=4, value=clock_in.clock_in_time.strftime("%H:%M") if clock_in.clock_in_time else "").border = thin_border
        ws.cell(row=row_num, column=5, value=clock_in.get_status_display()).border = thin_border
        ws.cell(row=row_num, column=6, value=clock_in.note).border = thin_border
        ws.cell(row=row_num, column=7, value=clock_in.recorded_by.get_full_name() if clock_in.recorded_by else "").border = thin_border
        ws.cell(row=row_num, column=8, value=str(clock_in.latitude) if clock_in.latitude else "").border = thin_border
        ws.cell(row=row_num, column=9, value=str(clock_in.longitude) if clock_in.longitude else "").border = thin_border
        ws.cell(row=row_num, column=10, value=str(clock_in.distance_from_school) if clock_in.distance_from_school else "").border = thin_border

        # Within radius cell with color
        within_cell = ws.cell(row=row_num, column=11, value="Yes" if clock_in.within_radius else ("No" if clock_in.within_radius is False else ""))
        within_cell.border = thin_border
        if clock_in.within_radius is True:
            within_cell.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")  # Green
            within_cell.font = Font(color="006100")
        elif clock_in.within_radius is False:
            within_cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")  # Red
            within_cell.font = Font(color="9C0006")
            pass

        row_num += 1
        pass

    # Auto-adjust column widths
    for col_num in range(1, len(headers) + 1):
        column_letter = get_column_letter(col_num)
        max_length = len(headers[col_num - 1])
        for cell in ws[column_letter]:
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))
                pass
            pass
        ws.column_dimensions[column_letter].width = min(max_length + 2, 50)
        pass

    # Create HTTP response
    response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    filename = f"clock_ins_{school.slug}_{datetime.date.today().strftime('%Y%m%d')}.xlsx"
    response["Content-Disposition"] = f'attachment; filename="{filename}"'

    wb.save(response)

    log_activity(request, "other", f"Exported {clock_ins.count()} clock-in records to Excel")

    return response
    pass
