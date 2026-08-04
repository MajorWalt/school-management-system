"""Excel export utilities for school reports."""

import io
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import datetime


def _get_header_style():
    """Get header styling for Excel sheets."""
    return {
        "font": Font(name="Arial", bold=True, color="FFFFFF", size=11),
        "fill": PatternFill("solid", start_color="1e40af"),
        "alignment": Alignment(horizontal="center", vertical="center", wrap_text=True),
    }


def _apply_style(cell, style):
    """Apply a style dictionary to a cell."""
    if "font" in style:
        cell.font = style["font"]
    if "fill" in style:
        cell.fill = style["fill"]
    if "alignment" in style:
        cell.alignment = style["alignment"]


def export_student_roster_excel(school, students, selected_form=None, selected_homeroom=None, status="enrolled"):
    """
    Generate Excel export for student roster.

    Args:
        school: School object
        students: Queryset or list of students
        selected_form: Optional selected Form
        selected_homeroom: Optional selected Homeroom
        status: Filter status applied

    Returns:
        HttpResponse with Excel file
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Student Roster"

    header_style = _get_header_style()

    # Title
    ws.merge_cells("A1:G1")
    title = ws["A1"]
    title.value = f"{school.name} - Student Roster"
    title.font = Font(name="Arial", bold=True, size=14)
    title.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 25

    # Filters info
    row = 2
    filters = []
    if selected_form:
        filters.append(f"Form: {selected_form.name}")
    if selected_homeroom:
        filters.append(f"Homeroom: {selected_homeroom.name}")
    if status != "all":
        filters.append(f"Status: {status.title()}")

    if filters:
        ws[f"A{row}"] = " | ".join(filters)
        ws[f"A{row}"].font = Font(size=10, italic=True)

    # Headers
    row = 4
    headers = ["Student ID", "Full Name", "Form", "Homeroom", "House", "Status", "Date Generated"]
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col)
        cell.value = header
        _apply_style(cell, header_style)

    # Column widths
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 12
    ws.column_dimensions["E"].width = 12
    ws.column_dimensions["F"].width = 14
    ws.column_dimensions["G"].width = 18

    # Data rows
    row = 5
    generated_at = datetime.datetime.now().strftime("%b %d, %Y %H:%M")

    for student in students:
        ws.cell(row=row, column=1).value = student.student_id
        ws.cell(row=row, column=2).value = student.get_full_name()
        ws.cell(row=row, column=3).value = student.form.name if student.form else ""
        ws.cell(row=row, column=4).value = student.homeroom.name if student.homeroom else ""
        ws.cell(row=row, column=5).value = student.house.name if student.house else ""
        ws.cell(row=row, column=6).value = student.current_status().title()
        ws.cell(row=row, column=7).value = generated_at
        row += 1

    # Save to buffer
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    # Create response
    filename = "student_roster.xlsx"
    response = HttpResponse(buf.read(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


def export_staff_list_excel(school, staff_list):
    """
    Generate Excel export for staff list.

    Args:
        school: School object
        staff_list: Queryset of Staff

    Returns:
        HttpResponse with Excel file
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Staff List"

    header_style = _get_header_style()

    # Title
    ws.merge_cells("A1:F1")
    title = ws["A1"]
    title.value = f"{school.name} - Staff Directory"
    title.font = Font(name="Arial", bold=True, size=14)
    title.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 25

    # Headers
    row = 3
    headers = ["Name", "Position", "Department", "Email", "Phone", "Status"]
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col)
        cell.value = header
        _apply_style(cell, header_style)

    # Column widths
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 25
    ws.column_dimensions["E"].width = 15
    ws.column_dimensions["F"].width = 12

    # Data rows
    row = 4
    for staff in staff_list:
        ws.cell(row=row, column=1).value = staff.get_full_name()
        ws.cell(row=row, column=2).value = staff.position or ""
        ws.cell(row=row, column=3).value = staff.department or ""
        ws.cell(row=row, column=4).value = staff.email or ""
        ws.cell(row=row, column=5).value = staff.phone or ""
        ws.cell(row=row, column=6).value = "Active" if staff.active else "Inactive"
        row += 1

    # Save to buffer
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    # Create response
    filename = "staff_list.xlsx"
    response = HttpResponse(buf.read(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


def export_class_list_excel(school, homeroom, students):
    """
    Generate Excel export for class list.

    Args:
        school: School object
        homeroom: Homeroom object
        students: List of students in homeroom

    Returns:
        HttpResponse with Excel file
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Class List"

    header_style = _get_header_style()

    # Title
    ws.merge_cells("A1:E1")
    title = ws["A1"]
    title.value = f"{school.name} - {homeroom.name} Class List"
    title.font = Font(name="Arial", bold=True, size=14)
    title.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 25

    # Subheader
    row = 2
    ws[f"A{row}"] = f"Form: {homeroom.form.name} | Total: {len(students)} students"
    ws[f"A{row}"].font = Font(size=10, italic=True)

    # Headers
    row = 4
    headers = ["#", "Student ID", "Full Name", "Date of Birth", "House"]
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col)
        cell.value = header
        _apply_style(cell, header_style)

    # Column widths
    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 22
    ws.column_dimensions["D"].width = 16
    ws.column_dimensions["E"].width = 14

    # Data rows
    row = 5
    for index, student in enumerate(students, start=1):
        ws.cell(row=row, column=1).value = index
        ws.cell(row=row, column=2).value = student.student_id
        ws.cell(row=row, column=3).value = student.get_full_name()
        ws.cell(row=row, column=4).value = student.date_of_birth.strftime("%b %d, %Y") if student.date_of_birth else ""
        ws.cell(row=row, column=5).value = student.house.name if student.house else ""
        row += 1

    # Save to buffer
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    # Create response
    filename = f"classlist_{homeroom.name.replace('/', '_')}.xlsx"
    response = HttpResponse(buf.read(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


def export_course_list_excel(school, section, enrolments):
    """
    Generate Excel export for course/section enrolment list.

    Args:
        school: School object
        section: Section object
        enrolments: List of enrolments

    Returns:
        HttpResponse with Excel file
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Course Enrolment"

    header_style = _get_header_style()

    # Title
    ws.merge_cells("A1:D1")
    title = ws["A1"]
    title.value = f"{school.name} - {section.course.name} Enrolment"
    title.font = Font(name="Arial", bold=True, size=14)
    title.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 25

    # Section info
    row = 2
    info_parts = [
        f"Form: {section.form.name}",
        f"Year: {section.academic_year.name}",
        f"Term: {section.term_number}",
    ]
    if section.teacher:
        info_parts.append(f"Teacher: {section.teacher.get_full_name()}")

    ws[f"A{row}"] = " | ".join(info_parts)
    ws[f"A{row}"].font = Font(size=10, italic=True)

    # Headers
    row = 4
    headers = ["#", "Student ID", "Full Name", "Homeroom"]
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col)
        cell.value = header
        _apply_style(cell, header_style)

    # Column widths
    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 22
    ws.column_dimensions["D"].width = 14

    # Data rows
    row = 5
    for index, enrolment in enumerate(enrolments, start=1):
        student = enrolment.student
        ws.cell(row=row, column=1).value = index
        ws.cell(row=row, column=2).value = student.student_id
        ws.cell(row=row, column=3).value = student.get_full_name()
        ws.cell(row=row, column=4).value = student.homeroom.name if student.homeroom else ""
        row += 1

    # Save to buffer
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    # Create response
    filename = f"courselist_{section.course.code or section.course.name}.xlsx"
    response = HttpResponse(buf.read(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


def export_attendance_summary_excel(school, homeroom_groups, grand, month_name, days_open):
    """
    Generate Excel export for attendance summary.

    Args:
        school: School object
        homeroom_groups: Dict of homeroom -> list of student records
        grand: Grand total dict with summary stats
        month_name: Month name (e.g., "January 2024")
        days_open: Number of school days

    Returns:
        HttpResponse with Excel file
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Attendance Summary"

    header_style = _get_header_style()

    # Title
    ws.merge_cells("A1:K1")
    title = ws["A1"]
    title.value = f"{school.name} - Attendance Summary {month_name}"
    title.font = Font(name="Arial", bold=True, size=14)
    title.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 25

    # Headers
    row = 3
    headers = [
        "Student",
        "Form/HR",
        "Days Enrolled",
        "Days Attended",
        "Absent (Total)",
        "Absent (Unexc.)",
        "Absent (Exc.)",
        "Absent (Other)",
        "Late (Unexc.)",
        "Late (Exc.)",
        "Att. %",
    ]
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col)
        cell.value = header
        _apply_style(cell, header_style)

    # Column widths
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 14
    ws.column_dimensions["E"].width = 14
    ws.column_dimensions["F"].width = 14
    ws.column_dimensions["G"].width = 14
    ws.column_dimensions["H"].width = 14
    ws.column_dimensions["I"].width = 14
    ws.column_dimensions["J"].width = 14
    ws.column_dimensions["K"].width = 10

    # Data rows
    row = 4
    for homeroom_name in sorted(homeroom_groups.keys()):
        records = homeroom_groups[homeroom_name]

        # Homeroom header
        ws.merge_cells(f"A{row}:K{row}")
        hr_cell = ws[f"A{row}"]
        hr_cell.value = homeroom_name
        hr_cell.font = Font(name="Arial", bold=True, size=11)
        hr_cell.fill = PatternFill("solid", start_color="E8E8E8")
        row += 1

        for record in records:
            ws.cell(row=row, column=1).value = record["student"].get_full_name()
            ws.cell(row=row, column=2).value = record["grade_homeroom"]
            ws.cell(row=row, column=3).value = record["enrolled"]
            ws.cell(row=row, column=4).value = record["attended"]
            ws.cell(row=row, column=5).value = record["absent_total"]
            ws.cell(row=row, column=6).value = record["absent_unexec"]
            ws.cell(row=row, column=7).value = record["absent_excused"]
            ws.cell(row=row, column=8).value = record["absent_other"]
            ws.cell(row=row, column=9).value = record["late_unexec"]
            ws.cell(row=row, column=10).value = record["late_excused"]
            ws.cell(row=row, column=11).value = f"{record['att_pct']}%"
            row += 1

    # Grand total row
    ws.merge_cells(f"A{row}:B{row}")
    grand_cell = ws[f"A{row}"]
    grand_cell.value = "GRAND TOTAL"
    grand_cell.font = Font(name="Arial", bold=True, size=11)
    grand_cell.fill = PatternFill("solid", start_color="D3D3D3")

    ws.cell(row=row, column=3).value = grand["enrolled"]
    ws.cell(row=row, column=3).font = Font(bold=True)
    ws.cell(row=row, column=4).value = grand["attended"]
    ws.cell(row=row, column=4).font = Font(bold=True)
    ws.cell(row=row, column=5).value = grand["absent_total"]
    ws.cell(row=row, column=5).font = Font(bold=True)
    ws.cell(row=row, column=6).value = grand["absent_unexec"]
    ws.cell(row=row, column=6).font = Font(bold=True)
    ws.cell(row=row, column=7).value = grand["absent_excused"]
    ws.cell(row=row, column=7).font = Font(bold=True)
    ws.cell(row=row, column=8).value = grand["absent_other"]
    ws.cell(row=row, column=8).font = Font(bold=True)
    ws.cell(row=row, column=9).value = grand["late_unexec"]
    ws.cell(row=row, column=9).font = Font(bold=True)
    ws.cell(row=row, column=10).value = grand["late_excused"]
    ws.cell(row=row, column=10).font = Font(bold=True)
    ws.cell(row=row, column=11).value = f"{grand['att_pct']}%"
    ws.cell(row=row, column=11).font = Font(bold=True)

    # Save to buffer
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    # Create response
    filename = f"attendance_{month_name.replace(' ', '_')}.xlsx"
    response = HttpResponse(buf.read(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


def export_merit_demerit_excel(school, groups, month_name, report_type, grand_points, grand_students):
    """
    Generate Excel export for merit/demerit report.

    Args:
        school: School object
        groups: Dict of homeroom -> list of student records with points
        month_name: Month name (e.g., "January 2024")
        report_type: "merit" or "demerit"
        grand_points: Total points across all students
        grand_students: Total number of students with records

    Returns:
        HttpResponse with Excel file
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Merit Demerit"

    header_style = _get_header_style()

    # Title
    report_label = "Merits" if report_type == "merit" else "Demerits"
    ws.merge_cells("A1:D1")
    title = ws["A1"]
    title.value = f"{school.name} - {report_label} Report {month_name}"
    title.font = Font(name="Arial", bold=True, size=14)
    title.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 25

    # Headers
    row = 3
    headers = ["Student", "Form/HR", "Points", "Record Count"]
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col)
        cell.value = header
        _apply_style(cell, header_style)

    # Column widths
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 14

    # Data rows
    row = 4
    for homeroom_name in sorted(groups.keys()):
        records = groups[homeroom_name]

        # Homeroom header
        ws.merge_cells(f"A{row}:D{row}")
        hr_cell = ws[f"A{row}"]
        hr_cell.value = homeroom_name
        hr_cell.font = Font(name="Arial", bold=True, size=11)
        hr_cell.fill = PatternFill("solid", start_color="E8E8E8")
        row += 1

        for record in records:
            ws.cell(row=row, column=1).value = record["student"].get_full_name()
            ws.cell(row=row, column=2).value = f"{record['student'].form.name if record['student'].form else ''}/"
            ws.cell(row=row, column=2).value += f"{record['student'].homeroom.name if record['student'].homeroom else ''}"
            ws.cell(row=row, column=3).value = record["points"]
            ws.cell(row=row, column=4).value = len(record["records"])
            row += 1

    # Summary section
    row += 1
    ws[f"A{row}"] = "Summary"
    ws[f"A{row}"].font = Font(bold=True, size=11)
    row += 1

    ws[f"A{row}"] = "Total Points:"
    ws[f"B{row}"] = grand_points
    ws[f"B{row}"].font = Font(bold=True)
    row += 1

    ws[f"A{row}"] = "Students with Records:"
    ws[f"B{row}"] = grand_students
    ws[f"B{row}"].font = Font(bold=True)

    # Save to buffer
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    # Create response
    filename = f"{'merits' if report_type == 'merit' else 'demerits'}_{month_name.replace(' ', '_')}.xlsx"
    response = HttpResponse(buf.read(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


def export_grade_report_excel(
    school, report_type, selected_section=None, selected_student=None, rows=None, evaluations=None, section=None, month_name=None, student=None
):
    """
    Generate Excel export for grade reports (by course, by student, gradebook, overview).

    Args:
        school: School object
        report_type: "by_course", "by_student", "gradebook", or "overview"
        selected_section: Section object (for by_course, gradebook)
        selected_student: Student object (for by_student)
        rows: List of student rows with grades (for by_course, gradebook)
        evaluations: List of evaluations (for by_course, gradebook)
        section: Section object (alternative param name for by_course)
        month_name: Month name (for gradebook)
        student: Student object (for by_student)

    Returns:
        HttpResponse with Excel file
    """
    # Handle param aliases
    if section and not selected_section:
        selected_section = section
    if student and not selected_student:
        selected_student = student

    wb = Workbook()
    ws = wb.active

    header_style = _get_header_style()

    if report_type == "by_course":
        ws.title = "Grade by Course"

        # Title
        ws.merge_cells("A1:E1")
        title = ws["A1"]
        title.value = f"{school.name} - Grades by Course"
        title.font = Font(name="Arial", bold=True, size=14)
        title.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 25

        if selected_section:
            # Course info
            row = 2
            ws[f"A{row}"] = f"{selected_section.course.name} ({selected_section.form.name})"
            ws[f"A{row}"].font = Font(size=10, italic=True)

            # Headers
            row = 4
            headers = ["Student", "Homeroom"] + [f"{ev.name}" for ev in evaluations]
            for col, header in enumerate(headers, start=1):
                cell = ws.cell(row=row, column=col)
                cell.value = header
                _apply_style(cell, header_style)

            # Column widths
            ws.column_dimensions["A"].width = 22
            ws.column_dimensions["B"].width = 14

            # Data rows
            row = 5
            for student_row in rows:
                ws.cell(row=row, column=1).value = student_row["student"].get_full_name()
                ws.cell(row=row, column=2).value = student_row["student"].homeroom.name if student_row["student"].homeroom else ""

                for col_idx, cell_data in enumerate(student_row["cells"], start=3):
                    if cell_data["absent"]:
                        ws.cell(row=row, column=col_idx).value = "ABS"
                    elif cell_data["pct"] is not None:
                        ws.cell(row=row, column=col_idx).value = f"{cell_data['pct']}%"
                    else:
                        ws.cell(row=row, column=col_idx).value = ""
                row += 1

    elif report_type == "by_student":
        ws.title = "Grade by Student"

        # Title
        ws.merge_cells("A1:D1")
        title = ws["A1"]
        title.value = f"{school.name} - Grades by Student"
        title.font = Font(name="Arial", bold=True, size=14)
        title.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 25

        if selected_student:
            # Student info
            row = 2
            ws[f"A{row}"] = f"{selected_student.get_full_name()}"
            ws[f"A{row}"].font = Font(size=10, italic=True)

            # Headers
            row = 4
            headers = ["Course", "Form", "Section", "Grade %"]
            for col, header in enumerate(headers, start=1):
                cell = ws.cell(row=row, column=col)
                cell.value = header
                _apply_style(cell, header_style)

            # Column widths
            ws.column_dimensions["A"].width = 22
            ws.column_dimensions["B"].width = 14
            ws.column_dimensions["C"].width = 14
            ws.column_dimensions["D"].width = 12

            # Data rows
            row = 5
            if rows:
                for student_row in rows:
                    ws.cell(row=row, column=1).value = student_row.get("course_name", "")
                    ws.cell(row=row, column=2).value = student_row.get("form_name", "")
                    ws.cell(row=row, column=3).value = student_row.get("section_name", "")
                    ws.cell(row=row, column=4).value = student_row.get("average", "")
                    row += 1

    elif report_type == "gradebook":
        ws.title = "Gradebook"

        # Title
        ws.merge_cells("A1:E1")
        title = ws["A1"]
        title.value = f"{school.name} - Gradebook"
        title.font = Font(name="Arial", bold=True, size=14)
        title.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 25

        if selected_section and month_name:
            # Section info
            row = 2
            ws[f"A{row}"] = f"{selected_section.course.name} ({selected_section.form.name}) - {month_name}"
            ws[f"A{row}"].font = Font(size=10, italic=True)

            # Headers
            row = 4
            headers = ["Student", "Homeroom"] + [f"{ev.name}" for ev in evaluations]
            for col, header in enumerate(headers, start=1):
                cell = ws.cell(row=row, column=col)
                cell.value = header
                _apply_style(cell, header_style)

            # Column widths
            ws.column_dimensions["A"].width = 22
            ws.column_dimensions["B"].width = 14

            # Data rows
            row = 5
            for student_row in rows:
                ws.cell(row=row, column=1).value = student_row["student"].get_full_name()
                ws.cell(row=row, column=2).value = student_row["student"].homeroom.name if student_row["student"].homeroom else ""

                for col_idx, cell_data in enumerate(student_row["cells"], start=3):
                    if cell_data["absent"]:
                        ws.cell(row=row, column=col_idx).value = "ABS"
                    elif cell_data["pct"] is not None:
                        ws.cell(row=row, column=col_idx).value = f"{cell_data['pct']}%"
                    else:
                        ws.cell(row=row, column=col_idx).value = ""
                row += 1

    else:  # overview
        ws.title = "Grade Overview"

        # Title
        ws.merge_cells("A1:E1")
        title = ws["A1"]
        title.value = f"{school.name} - Grade Overview"
        title.font = Font(name="Arial", bold=True, size=14)
        title.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 25

        # Headers
        row = 3
        headers = ["Form", "Course", "Section", "Avg. Grade", "Students"]
        for col, header in enumerate(headers, start=1):
            cell = ws.cell(row=row, column=col)
            cell.value = header
            _apply_style(cell, header_style)

        # Column widths
        ws.column_dimensions["A"].width = 14
        ws.column_dimensions["B"].width = 22
        ws.column_dimensions["C"].width = 14
        ws.column_dimensions["D"].width = 12
        ws.column_dimensions["E"].width = 12

        # Data rows
        row = 4
        if rows:
            for section_row in rows:
                ws.cell(row=row, column=1).value = section_row.get("form", "")
                ws.cell(row=row, column=2).value = section_row.get("course", "")
                ws.cell(row=row, column=3).value = section_row.get("section", "")
                ws.cell(row=row, column=4).value = section_row.get("avg", "")
                ws.cell(row=row, column=5).value = section_row.get("count", "")
                row += 1

    # Save to buffer
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    # Create response
    filename = f"grades_{report_type}.xlsx"
    response = HttpResponse(buf.read(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response
